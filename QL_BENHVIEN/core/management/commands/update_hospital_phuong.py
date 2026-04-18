import csv
import os
import unicodedata

from django.core.management.base import BaseCommand, CommandError

from core.models import BenhVien


def _normalize_text(value):
    text = str(value or "").strip().lower()
    if not text:
        return ""
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.replace("_", " ").replace("-", " ")
    return " ".join(text.split())


def _map_headers(headers):
    aliases = {
        "ten": {
            "ten",
            "ten benh vien",
            "benh vien",
            "hospital",
            "name",
        },
        "phuong": {
            "phuong",
            "phuong xa",
            "phuong/xa",
            "ward",
        },
    }
    normalized = [_normalize_text(h) for h in headers]
    mapped = {}
    for idx, header in enumerate(normalized):
        for key, options in aliases.items():
            if header in options:
                mapped[key] = idx
                break
    return mapped


def _read_csv_rows(path):
    for encoding in ("utf-8-sig", "utf-8", "cp1258"):
        try:
            with open(path, "r", encoding=encoding, newline="") as handle:
                rows = list(csv.reader(handle))
            return rows
        except UnicodeDecodeError:
            continue
    raise CommandError("Khong doc duoc file CSV. Vui long dung UTF-8.")


def _read_xlsx_rows(path, sheet_name=None):
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise CommandError(
            "Missing dependency openpyxl. Install with: pip install openpyxl"
        ) from exc

    wb = load_workbook(path, data_only=True, read_only=True)
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise CommandError(f"Khong tim thay sheet '{sheet_name}'.")
        ws = wb[sheet_name]
    else:
        ws = wb.active
    return [list(row) for row in ws.iter_rows(values_only=True)]


class Command(BaseCommand):
    help = "Cap nhat phuong cho benh vien theo mapping ten benh vien -> phuong."

    def add_arguments(self, parser):
        parser.add_argument("file", help="Duong dan file mapping (.csv hoac .xlsx)")
        parser.add_argument("--sheet", default=None, help="Ten sheet neu dung xlsx")
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Chi xem truoc, khong ghi vao database.",
        )

    def handle(self, *args, **options):
        path = options["file"]
        sheet_name = options["sheet"]
        dry_run = options["dry_run"]

        if not os.path.exists(path):
            raise CommandError(f"Khong tim thay file: {path}")

        ext = os.path.splitext(path)[1].lower()
        if ext == ".csv":
            rows = _read_csv_rows(path)
        elif ext in {".xlsx", ".xlsm"}:
            rows = _read_xlsx_rows(path, sheet_name=sheet_name)
        else:
            raise CommandError("Chi ho tro .csv hoac .xlsx/.xlsm")

        if not rows:
            raise CommandError("File mapping trong.")

        header_map = _map_headers(rows[0])
        if "ten" not in header_map or "phuong" not in header_map:
            raise CommandError(
                "File can co cot ten benh vien va phuong (vi du: ten, phuong)."
            )

        hospitals_by_name = {}
        for hospital in BenhVien.objects.all():
            key = _normalize_text(hospital.ten)
            hospitals_by_name.setdefault(key, []).append(hospital)

        updated = 0
        unchanged = 0
        not_found = 0
        skipped = 0
        duplicate_name_hits = 0

        for row in rows[1:]:
            ten = row[header_map["ten"]] if header_map["ten"] < len(row) else None
            phuong = row[header_map["phuong"]] if header_map["phuong"] < len(row) else None

            ten_text = str(ten or "").strip()
            phuong_text = str(phuong or "").strip()
            if not ten_text or not phuong_text:
                skipped += 1
                continue

            matches = hospitals_by_name.get(_normalize_text(ten_text), [])
            if not matches:
                not_found += 1
                self.stdout.write(self.style.WARNING(f"Khong tim thay: {ten_text}"))
                continue

            if len(matches) > 1:
                duplicate_name_hits += 1

            for hospital in matches:
                if hospital.phuong == phuong_text:
                    unchanged += 1
                    continue
                if not dry_run:
                    hospital.phuong = phuong_text
                    hospital.save(update_fields=["phuong"])
                updated += 1

        mode = "DRY-RUN" if dry_run else "UPDATED"
        self.stdout.write(self.style.SUCCESS(f"[{mode}] Hoan tat cap nhat phuong."))
        self.stdout.write(f"- Updated: {updated}")
        self.stdout.write(f"- Unchanged: {unchanged}")
        self.stdout.write(f"- Not found: {not_found}")
        self.stdout.write(f"- Skipped (thieu du lieu): {skipped}")
        self.stdout.write(f"- Duplicate name groups hit: {duplicate_name_hits}")
