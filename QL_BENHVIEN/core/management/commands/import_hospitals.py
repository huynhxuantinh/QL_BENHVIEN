import datetime
import unicodedata
from pathlib import Path

from django.contrib.gis.geos import Point
from django.core.files import File
from django.core.management.base import BaseCommand, CommandError
from django.db.models import Max

from core.models import (
    BacSi,
    BenhVien,
    BenhVienHinhAnh,
    GioLamViecBacSi,
    GioLamViecBenhVien,
    Khoa,
)


def _normalize_header(value):
    text = str(value).strip().lower()
    if not text:
        return ""
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    # Normalize common separators so headers like "benh_vien" and "benh-vien"
    # are treated the same as "benh vien".
    text = text.replace("_", " ").replace("-", " ")
    text = text.replace("\n", " ").replace("\r", " ")
    return " ".join(text.split())


def _parse_bool(value, default=False):
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    text = _normalize_header(value)
    return text in {"1", "true", "yes", "y", "co", "on"}


def _parse_time(value, default):
    if value is None or value == "":
        return default
    if isinstance(value, datetime.time):
        return value
    if isinstance(value, datetime.datetime):
        return value.time()
    text = str(value).strip()
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    return default


def _parse_day(value):
    if value is None or value == "":
        return None
    if isinstance(value, int):
        if 0 <= value <= 6:
            return value
        if 1 <= value <= 7:
            return 0 if value == 7 else value
    text = _normalize_header(value)
    if not text:
        return None
    if text in {"cn", "chu nhat", "sun", "sunday"}:
        return 0
    for digit in ("2", "3", "4", "5", "6", "7"):
        if digit in text:
            return int(digit) - 1
    return None


def _parse_loai_hinh(value):
    if not value:
        return "cong"
    text = _normalize_header(value)
    if text in {"cong", "cong lap", "benh vien cong"}:
        return "cong"
    if text in {"tu", "tu nhan", "benh vien tu"}:
        return "tu"
    if text in {"qt", "quoc te", "quoc te"}:
        return "qt"
    return "cong"


def _map_headers(headers, aliases):
    normalized_headers = [_normalize_header(h) for h in headers]
    header_map = {}
    for idx, name in enumerate(normalized_headers):
        for key, values in aliases.items():
            if name in values:
                header_map[key] = idx
                break
    return header_map


def _find_sheet(workbook, *names):
    wanted = {_normalize_header(name) for name in names}
    for ws in workbook.worksheets:
        if _normalize_header(ws.title) in wanted:
            return ws
    return None


def _get_cell(row, header_map, key, default=None):
    if key not in header_map:
        return default
    return row[header_map[key]]


def _get_hospital_by_name(name, cache):
    if not name:
        return None
    key = _normalize_header(name)
    if key in cache:
        return cache[key]
    bv = BenhVien.objects.filter(ten__iexact=str(name).strip()).first()
    if bv:
        cache[key] = bv
    return bv


def _get_khoa_by_name(name, benh_vien, cache):
    if not name or not benh_vien:
        return None
    key = (_normalize_header(name), benh_vien.id)
    if key in cache:
        return cache[key]
    khoa = Khoa.objects.filter(ten__iexact=str(name).strip(), benh_vien=benh_vien).first()
    if khoa:
        cache[key] = khoa
    return khoa


def _split_image_paths(raw_value):
    if raw_value is None:
        return []
    text = str(raw_value).strip()
    if not text:
        return []
    # Support common separators in Excel cells.
    for sep in ("|", "\n", "\r"):
        text = text.replace(sep, ";")
    return [part.strip() for part in text.split(";") if part.strip()]


def _looks_like_image_header(normalized_header):
    if not normalized_header:
        return False
    if "folder" in normalized_header or "thu muc" in normalized_header:
        return False
    if normalized_header in {
        "hinh anh",
        "duong dan anh",
        "image",
        "image path",
        "image url",
        "anh",
    }:
        return True
    prefixes = (
        "hinh anh ",
        "duong dan anh ",
        "image ",
        "image path ",
        "image url ",
        "anh ",
    )
    return any(normalized_header.startswith(prefix) for prefix in prefixes)


def _collect_image_paths(row, header_names, source_dir):
    raw_paths = []
    for idx, header_name in enumerate(header_names):
        if _looks_like_image_header(header_name):
            raw_paths.extend(_split_image_paths(row[idx] if idx < len(row) else None))

    resolved = []
    for raw_path in raw_paths:
        candidate = Path(raw_path)
        if candidate.is_absolute():
            resolved.append(candidate.resolve())
            continue

        # Try common bases: file folder, project root (parent), current working dir.
        base_candidates = [source_dir, source_dir.parent, Path.cwd()]
        picked = None
        for base in base_candidates:
            trial = (base / candidate).resolve()
            if trial.exists():
                picked = trial
                break
        resolved.append(picked or (source_dir / candidate).resolve())
    return resolved


def _import_hospital_images(benh_vien, image_paths):
    if not image_paths:
        return {"created": 0, "skipped": 0}

    created = 0
    skipped = 0
    has_cover = benh_vien.hinh_anhs.filter(la_anh_dai_dien=True).exists()
    next_order = benh_vien.hinh_anhs.aggregate(max_order=Max("thu_tu")).get("max_order")
    next_order = 0 if next_order is None else int(next_order) + 1

    for image_path in image_paths:
        if not image_path.exists() or not image_path.is_file():
            skipped += 1
            continue
        with image_path.open("rb") as f:
            image = BenhVienHinhAnh(
                benh_vien=benh_vien,
                thu_tu=next_order,
                la_anh_dai_dien=not has_cover and created == 0,
            )
            image.hinh_anh.save(image_path.name, File(f), save=False)
            image.save()
        created += 1
        next_order += 1
    return {"created": created, "skipped": skipped}


def _import_hospitals(ws, update_existing, source_dir):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"created": 0, "updated": 0, "skipped": 0, "images_created": 0, "images_skipped": 0}

    aliases = {
        "ten": {"ten", "ten benh vien", "name", "hospital", "benh vien"},
        "dia_chi": {"dia chi", "diachi", "address"},
        "phuong": {"phuong", "phuong/xa", "ward", "quan", "quan/huyen", "district"},
        "lat": {"lat", "latitude", "vi do", "vido", "y"},
        "lon": {"lon", "lng", "longitude", "kinh do", "kinhdo", "x"},
        "co_bhyt": {"co bhyt", "bhyt", "bao hiem y te", "bhyt?"},
        "cap_cuu_24h": {"cap cuu 24h", "cap cuu 24", "capcuu 24h"},
        "loai_hinh": {"loai hinh", "loai", "type"},
        "gio_mo": {"gio mo", "open", "open time"},
        "gio_dong": {"gio dong", "close", "close time"},
    }

    header_map = _map_headers(rows[0], aliases)
    header_names = [_normalize_header(h) for h in rows[0]]
    missing_required = [k for k in ("ten", "lat", "lon") if k not in header_map]
    if missing_required:
        raise CommandError(f"Missing required columns in Hospitals sheet: {', '.join(missing_required)}")

    created = 0
    updated = 0
    skipped = 0
    images_created = 0
    images_skipped = 0

    default_open = datetime.time(7, 0)
    default_close = datetime.time(17, 0)

    for row in rows[1:]:
        ten = _get_cell(row, header_map, "ten")
        if not ten:
            skipped += 1
            continue

        lat = _get_cell(row, header_map, "lat")
        lon = _get_cell(row, header_map, "lon")
        try:
            lat = float(lat)
            lon = float(lon)
        except (TypeError, ValueError):
            skipped += 1
            continue

        if not (-90 <= lat <= 90 and -180 <= lon <= 180):
            skipped += 1
            continue

        dia_chi = _get_cell(row, header_map, "dia_chi", "") or ""
        phuong = _get_cell(row, header_map, "phuong", "Chua ro") or "Chua ro"

        gio_mo = _parse_time(_get_cell(row, header_map, "gio_mo"), default_open)
        gio_dong = _parse_time(_get_cell(row, header_map, "gio_dong"), default_close)

        co_bhyt = _parse_bool(_get_cell(row, header_map, "co_bhyt"), default=False)
        cap_cuu_24h = _parse_bool(_get_cell(row, header_map, "cap_cuu_24h"), default=False)
        loai_hinh = _parse_loai_hinh(_get_cell(row, header_map, "loai_hinh"))

        bv = None
        if update_existing:
            by_name_qs = BenhVien.objects.filter(ten__iexact=str(ten).strip()).order_by("id")
            if dia_chi:
                exact_qs = by_name_qs.filter(dia_chi__iexact=str(dia_chi).strip())
                bv = exact_qs.first()
            if not bv:
                for candidate in by_name_qs:
                    if not candidate.vi_tri:
                        continue
                    if abs(candidate.vi_tri.x - lon) < 1e-6 and abs(candidate.vi_tri.y - lat) < 1e-6:
                        bv = candidate
                        break
            if not bv:
                # Fallback: update by name even when address changed (e.g. renamed wards).
                bv = by_name_qs.first()

        if bv:
            bv.dia_chi = dia_chi
            bv.phuong = phuong
            bv.vi_tri = Point(lon, lat, srid=4326)
            bv.cap_cuu_24h = cap_cuu_24h
            bv.co_bhyt = co_bhyt
            bv.gio_mo = gio_mo
            bv.gio_dong = gio_dong
            bv.loai_hinh = loai_hinh
            bv.save()
            updated += 1
        else:
            bv = BenhVien.objects.create(
                ten=ten,
                dia_chi=dia_chi,
                phuong=phuong,
                vi_tri=Point(lon, lat, srid=4326),
                cap_cuu_24h=cap_cuu_24h,
                co_bhyt=co_bhyt,
                gio_mo=gio_mo,
                gio_dong=gio_dong,
                loai_hinh=loai_hinh,
            )
            created += 1

        image_paths = _collect_image_paths(row, header_names, source_dir)
        image_result = _import_hospital_images(bv, image_paths)
        images_created += image_result["created"]
        images_skipped += image_result["skipped"]

        if not GioLamViecBenhVien.objects.filter(benh_vien=bv).exists():
            for thu in range(7):
                nghi = thu == 0
                GioLamViecBenhVien.objects.create(
                    benh_vien=bv,
                    thu=thu,
                    gio_mo=datetime.time(0, 0) if nghi else gio_mo,
                    gio_dong=datetime.time(0, 0) if nghi else gio_dong,
                    nghi=nghi,
                )

    return {
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "images_created": images_created,
        "images_skipped": images_skipped,
    }


def _import_departments(ws, benh_vien_cache):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"created": 0, "skipped": 0}

    aliases = {
        "ten": {"ten", "khoa", "ten khoa", "department"},
        "benh_vien": {"benh vien", "benhvien", "hospital", "ten benh vien"},
    }
    header_map = _map_headers(rows[0], aliases)
    missing_required = [k for k in ("ten", "benh_vien") if k not in header_map]
    if missing_required:
        raise CommandError(f"Missing required columns in Departments sheet: {', '.join(missing_required)}")

    created = 0
    skipped = 0
    khoa_cache = {}

    for row in rows[1:]:
        ten = _get_cell(row, header_map, "ten")
        benh_vien_name = _get_cell(row, header_map, "benh_vien")
        if not ten or not benh_vien_name:
            skipped += 1
            continue
        benh_vien = _get_hospital_by_name(benh_vien_name, benh_vien_cache)
        if not benh_vien:
            skipped += 1
            continue
        key = (_normalize_header(ten), benh_vien.id)
        if key in khoa_cache:
            continue
        khoa, was_created = Khoa.objects.get_or_create(
            ten=str(ten).strip(),
            benh_vien=benh_vien,
        )
        khoa_cache[key] = khoa
        if was_created:
            created += 1

    return {"created": created, "skipped": skipped}


def _import_doctors(ws, benh_vien_cache):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"created": 0, "updated": 0, "skipped": 0}

    aliases = {
        "ho_ten": {"ho ten", "ten", "bac si", "doctor"},
        "so_dien_thoai": {"so dien thoai", "sdt", "phone"},
        "chuyen_khoa": {"chuyen khoa", "chuyen mon", "specialty"},
        "khoa": {"khoa", "department"},
        "benh_vien": {"benh vien", "benhvien", "hospital"},
    }
    header_map = _map_headers(rows[0], aliases)
    missing_required = [k for k in ("ho_ten", "so_dien_thoai", "khoa", "benh_vien") if k not in header_map]
    if missing_required:
        raise CommandError(f"Missing required columns in Doctors sheet: {', '.join(missing_required)}")

    created = 0
    updated = 0
    skipped = 0

    khoa_cache = {}

    for row in rows[1:]:
        ho_ten = _get_cell(row, header_map, "ho_ten")
        so_dien_thoai = _get_cell(row, header_map, "so_dien_thoai")
        khoa_name = _get_cell(row, header_map, "khoa")
        benh_vien_name = _get_cell(row, header_map, "benh_vien")
        chuyen_khoa = _get_cell(row, header_map, "chuyen_khoa")

        if not ho_ten or not so_dien_thoai or not khoa_name or not benh_vien_name:
            skipped += 1
            continue

        benh_vien = _get_hospital_by_name(benh_vien_name, benh_vien_cache)
        if not benh_vien:
            skipped += 1
            continue

        khoa = _get_khoa_by_name(khoa_name, benh_vien, khoa_cache)
        if not khoa:
            khoa, _ = Khoa.objects.get_or_create(
                ten=str(khoa_name).strip(),
                benh_vien=benh_vien,
            )
            khoa_cache[(_normalize_header(khoa_name), benh_vien.id)] = khoa

        chuyen_khoa = str(chuyen_khoa).strip() if chuyen_khoa else str(khoa_name).strip()
        phone = str(so_dien_thoai).strip()

        if not phone:
            skipped += 1
            continue

        if BacSi.objects.filter(so_dien_thoai=phone).exists():
            bac_si = BacSi.objects.get(so_dien_thoai=phone)
            bac_si.ho_ten = str(ho_ten).strip()
            bac_si.chuyen_khoa = chuyen_khoa
            bac_si.khoa = khoa
            bac_si.benh_vien = benh_vien
            bac_si.save()
            updated += 1
        else:
            BacSi.objects.create(
                ho_ten=str(ho_ten).strip(),
                chuyen_khoa=chuyen_khoa,
                khoa=khoa,
                benh_vien=benh_vien,
                so_dien_thoai=phone,
            )
            created += 1

    return {"created": created, "updated": updated, "skipped": skipped}


def _import_doctor_schedules(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"created": 0, "updated": 0, "skipped": 0}

    aliases = {
        "bac_si": {"bac si", "doctor", "so dien thoai", "sdt"},
        "thu": {"thu", "day", "weekday"},
        "gio_bat_dau": {"gio bat dau", "start", "start time"},
        "gio_ket_thuc": {"gio ket thuc", "end", "end time"},
        "nghi": {"nghi", "off"},
    }
    header_map = _map_headers(rows[0], aliases)
    missing_required = [k for k in ("bac_si", "thu") if k not in header_map]
    if missing_required:
        raise CommandError(f"Missing required columns in DoctorSchedules sheet: {', '.join(missing_required)}")

    created = 0
    updated = 0
    skipped = 0

    default_start = datetime.time(7, 30)
    default_end = datetime.time(16, 30)

    for row in rows[1:]:
        bac_si_key = _get_cell(row, header_map, "bac_si")
        thu = _parse_day(_get_cell(row, header_map, "thu"))
        if not bac_si_key or thu is None:
            skipped += 1
            continue

        bac_si = BacSi.objects.filter(so_dien_thoai=str(bac_si_key).strip()).first()
        if not bac_si:
            bac_si = BacSi.objects.filter(ho_ten__iexact=str(bac_si_key).strip()).first()
        if not bac_si:
            skipped += 1
            continue

        nghi = _parse_bool(_get_cell(row, header_map, "nghi"), default=False)
        gio_bat_dau = _parse_time(_get_cell(row, header_map, "gio_bat_dau"), default_start)
        gio_ket_thuc = _parse_time(_get_cell(row, header_map, "gio_ket_thuc"), default_end)

        obj, was_created = GioLamViecBacSi.objects.update_or_create(
            bac_si=bac_si,
            thu=thu,
            defaults={
                "gio_bat_dau": datetime.time(0, 0) if nghi else gio_bat_dau,
                "gio_ket_thuc": datetime.time(0, 0) if nghi else gio_ket_thuc,
                "nghi": nghi,
            },
        )
        if was_created:
            created += 1
        else:
            updated += 1

    return {"created": created, "updated": updated, "skipped": skipped}


class Command(BaseCommand):
    help = "Import hospitals (and optional departments/doctors) from an Excel file (.xlsx)."

    def add_arguments(self, parser):
        parser.add_argument("file", nargs="?", help="Path to .xlsx file")
        parser.add_argument("--file", dest="file_opt", help="Path to .xlsx file")
        parser.add_argument("--sheet", default=None, help="Sheet name (optional)")
        parser.add_argument(
            "--update-existing",
            action="store_true",
            help="Update existing hospitals by name + address",
        )

    def handle(self, *args, **options):
        try:
            from openpyxl import load_workbook
        except Exception as exc:  # pragma: no cover - import error path
            raise CommandError(
                "Missing dependency openpyxl. Install with: pip install openpyxl"
            ) from exc

        path = options.get("file_opt") or options.get("file")
        if not path:
            raise CommandError("Missing file path. Provide a .xlsx path.")
        source_dir = Path(path).expanduser().resolve().parent
        sheet_name = options["sheet"]
        update_existing = options["update_existing"]

        wb = load_workbook(path, data_only=True)

        if sheet_name:
            ws = wb[sheet_name]
            result = _import_hospitals(ws, update_existing, source_dir)
            self.stdout.write(
                self.style.SUCCESS(
                    "Hospitals import finished. "
                    f"Created: {result['created']}, Updated: {result['updated']}, Skipped: {result['skipped']}. "
                    f"Images created: {result['images_created']}, Images skipped: {result['images_skipped']}."
                )
            )
            return

        hospitals_ws = _find_sheet(wb, "Hospitals", "BenhVien", "Benh Vien", "benh_vien")
        departments_ws = _find_sheet(wb, "Departments", "Khoa", "khoa")
        doctors_ws = _find_sheet(wb, "Doctors", "BacSi", "Bac Si", "bac_si")
        schedules_ws = _find_sheet(
            wb,
            "DoctorSchedules",
            "LichBacSi",
            "Lich Bac Si",
            "gio_lam_viec_bac_si",
        )

        if not any([hospitals_ws, departments_ws, doctors_ws, schedules_ws]):
            ws = wb.active
            result = _import_hospitals(ws, update_existing, source_dir)
            self.stdout.write(
                self.style.SUCCESS(
                    "Hospitals import finished. "
                    f"Created: {result['created']}, Updated: {result['updated']}, Skipped: {result['skipped']}. "
                    f"Images created: {result['images_created']}, Images skipped: {result['images_skipped']}."
                )
            )
            return

        if hospitals_ws:
            result = _import_hospitals(hospitals_ws, update_existing, source_dir)
            self.stdout.write(
                self.style.SUCCESS(
                    "Hospitals import finished. "
                    f"Created: {result['created']}, Updated: {result['updated']}, Skipped: {result['skipped']}. "
                    f"Images created: {result['images_created']}, Images skipped: {result['images_skipped']}."
                )
            )
        else:
            self.stdout.write(self.style.WARNING("Hospitals sheet not found. Skipping hospitals import."))

        benh_vien_cache = {_normalize_header(bv.ten): bv for bv in BenhVien.objects.all()}

        if departments_ws:
            try:
                dep_result = _import_departments(departments_ws, benh_vien_cache)
            except CommandError as exc:
                self.stdout.write(
                    self.style.WARNING(f"Departments sheet skipped: {exc}")
                )
            else:
                self.stdout.write(
                    self.style.SUCCESS(
                        "Departments import finished. "
                        f"Created: {dep_result['created']}, Skipped: {dep_result['skipped']}."
                    )
                )
        else:
            self.stdout.write(self.style.WARNING("Departments sheet not found. Skipping departments import."))

        if doctors_ws:
            try:
                doc_result = _import_doctors(doctors_ws, benh_vien_cache)
            except CommandError as exc:
                self.stdout.write(
                    self.style.WARNING(f"Doctors sheet skipped: {exc}")
                )
            else:
                self.stdout.write(
                    self.style.SUCCESS(
                        "Doctors import finished. "
                        f"Created: {doc_result['created']}, Updated: {doc_result['updated']}, Skipped: {doc_result['skipped']}."
                    )
                )
        else:
            self.stdout.write(self.style.WARNING("Doctors sheet not found. Skipping doctors import."))

        if schedules_ws:
            try:
                sch_result = _import_doctor_schedules(schedules_ws)
            except CommandError as exc:
                self.stdout.write(
                    self.style.WARNING(f"DoctorSchedules sheet skipped: {exc}")
                )
            else:
                self.stdout.write(
                    self.style.SUCCESS(
                        "Doctor schedules import finished. "
                        f"Created: {sch_result['created']}, Updated: {sch_result['updated']}, Skipped: {sch_result['skipped']}."
                    )
                )
        else:
            self.stdout.write(self.style.WARNING("DoctorSchedules sheet not found. Skipping schedules import."))
